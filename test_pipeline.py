import os
import sys
import json
import pymupdf
import docx
import openpyxl
from PIL import Image, ImageDraw, ImageFont

from document_reader import read_document

sys.stdout.reconfigure(encoding='utf-8')

TEST_DIR = r"d:\AI\ocr\test-ocr\test_samples"

def create_synthetic_test_files():
    os.makedirs(TEST_DIR, exist_ok=True)
    created_files = {}

    # 1. PDF Text & Vietnamese
    pdf_text_path = os.path.join(TEST_DIR, "1_pdf_text_vietnamese.pdf")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc", fontsize=14)
    page.insert_text((50, 100), "HỢP ĐỒNG MUA BÁN HÀNG HÓA", fontsize=16)
    page.insert_text((50, 150), "ĐIỀU 1. ĐỐI TƯỢNG HỢP ĐỒNG\nBên A đồng ý bán cho Bên B sản phẩm máy tính.", fontsize=12)
    doc.save(pdf_text_path)
    doc.close()
    created_files["1_pdf_text"] = pdf_text_path

    # 2. PDF Scan (Rendered text as image into PDF page)
    pdf_scan_path = os.path.join(TEST_DIR, "2_pdf_scan.pdf")
    img = Image.new('RGB', (600, 400), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((50, 50), "HÓA ĐƠN GIÁ TRỊ GIA TĂNG (SCAN)", fill=(0, 0, 0))
    d.text((50, 100), "Tên hàng: Máy in HP 2026", fill=(0, 0, 0))
    d.text((50, 150), "Thành tiền: 5.000.000 VNĐ", fill=(0, 0, 0))
    img_scan_path = os.path.join(TEST_DIR, "temp_scan.png")
    img.save(img_scan_path)

    doc_scan = pymupdf.open()
    rect = pymupdf.Rect(0, 0, 600, 400)
    page_scan = doc_scan.new_page(width=600, height=400)
    page_scan.insert_image(rect, filename=img_scan_path)
    doc_scan.save(pdf_scan_path)
    doc_scan.close()
    created_files["2_pdf_scan"] = pdf_scan_path

    # 3. PDF Mixed (Page 1 text, Page 2 scan)
    pdf_mixed_path = os.path.join(TEST_DIR, "3_pdf_mixed.pdf")
    doc_mixed = pymupdf.open()
    # Page 1: Native Text
    p1 = doc_mixed.new_page()
    p1.insert_text((50, 50), "BÁO CÁO TÀI CHÍNH NĂM 2026", fontsize=14)
    p1.insert_text((50, 100), "Trang 1 có văn bản kỹ thuật số chuẩn.", fontsize=12)
    # Page 2: Image Scan
    p2 = doc_mixed.new_page(width=600, height=400)
    p2.insert_image(rect, filename=img_scan_path)
    doc_mixed.save(pdf_mixed_path)
    doc_mixed.close()
    created_files["3_pdf_mixed"] = pdf_mixed_path

    # 4. DOCX with Table
    docx_path = os.path.join(TEST_DIR, "4_doc_with_table.docx")
    doc_docx = docx.Document()
    doc_docx.add_heading("HỢP ĐỒNG CUNG CẤP DỊCH VỤ", level=1)
    doc_docx.add_paragraph("ĐIỀU 1. DANH MỤC SẢN PHẨM")
    table = doc_docx.add_table(rows=2, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'STT'
    hdr_cells[1].text = 'Tên hàng'
    hdr_cells[2].text = 'Thành tiền'
    row_cells = table.rows[1].cells
    row_cells[0].text = '1'
    row_cells[1].text = 'Phần mềm OCR 2026'
    row_cells[2].text = '15.000.000 VNĐ'
    doc_docx.save(docx_path)
    created_files["4_docx"] = docx_path

    # 5. XLSX
    xlsx_path = os.path.join(TEST_DIR, "5_excel_sheet.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Báo_Cáo_Doanh_Thu"
    ws.append(["Mã HĐ", "Khách hàng", "Số tiền"])
    ws.append(["HD001", "Công ty ABC", 50000000])
    ws.append(["HD002", "Công ty XYZ", 75000000])
    wb.save(xlsx_path)
    created_files["5_xlsx"] = xlsx_path

    # 6. PNG Image
    img_png_path = os.path.join(TEST_DIR, "6_receipt_image.png")
    img_receipt = Image.new('RGB', (500, 300), color=(255, 255, 255))
    d_r = ImageDraw.Draw(img_receipt)
    d_r.text((30, 30), "BIÊN NHẬN CHUYỂN TIỀN THÀNH CÔNG", fill=(0, 0, 0))
    d_r.text((30, 80), "Số tiền: 2.500.000 VND", fill=(0, 0, 0))
    d_r.text((30, 130), "Tài khoản hưởng: 123456789 VIETCOMBANK", fill=(0, 0, 0))
    img_receipt.save(img_png_path)
    created_files["6_image"] = img_png_path

    # 7. PDF Text Fallback (Font corrupted / Garbage CID characters)
    pdf_corrupted_path = os.path.join(TEST_DIR, "7_pdf_corrupted_text.pdf")
    doc_corrupt = pymupdf.open()
    p_c = doc_corrupt.new_page()
    # Chèn các ký tự CID rác giả lập lỗi mã hóa font
    p_c.insert_text((50, 50), "(cid:12) (cid:45) (cid:89) (cid:100) (cid:101) (cid:102) (cid:103)", fontsize=12)
    doc_corrupt.save(pdf_corrupted_path)
    doc_corrupt.close()
    created_files["7_pdf_corrupted"] = pdf_corrupted_path

    return created_files

def run_pipeline_test():
    print("\n" + "=" * 70)
    print("        RUNNING COMPREHENSIVE PIPELINE TEST SUITE")
    print("=" * 70)

    files = create_synthetic_test_files()
    test_summary = []

    for name, path in files.items():
        print(f"\n------------------------------------------------------------")
        print(f"▶ Testing Scenario: [{name}] ({os.path.basename(path)})")
        print(f"------------------------------------------------------------")

        res = read_document(path)
        
        fallback_occurred = any(pg.get("method") == "ocr" for pg in res.get("pages", [])) if res.get("document_type") == "pdf" else (res.get("extraction_method") == "ocr")

        scenario_report = {
            "scenario": name,
            "filename": os.path.basename(path),
            "document_type": res.get("document_type"),
            "source_type": res.get("source_type"),
            "extraction_method": res.get("extraction_method"),
            "total_pages": res.get("total_pages"),
            "has_table": res.get("has_table"),
            "fallback_ocr": fallback_occurred,
            "pages_summary": [
                {"page": p.get("page"), "method": p.get("method"), "confidence": p.get("confidence")}
                for p in res.get("pages", [])
            ]
        }
        test_summary.append(scenario_report)

        print(f"  ✓ Source Type Selected:  {res.get('source_type')}")
        print(f"  ✓ Extraction Method:     {res.get('extraction_method')}")
        print(f"  ✓ Total Pages:           {res.get('total_pages')}")
        print(f"  ✓ Has Tables Detected:   {res.get('has_table')}")
        print(f"  ✓ Fallback OCR Triggered:{fallback_occurred}")
        print(f"  ✓ Extract Text Snippet:  {res.get('full_text', '')[:100].strip()}...")

    print("\n" + "=" * 70)
    print("                TEST SUITE SUMMARY REPORT")
    print("=" * 70)
    print(json.dumps(test_summary, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    run_pipeline_test()
